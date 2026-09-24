"""Convert the REC² workbook into the map's static data file.

Usage: python scripts/build_data.py materials/REC2_data_template.xlsx docs/data.json

The published file excludes email addresses by default. For an access-controlled
internal deployment only, add --include-internal-email.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

def clean(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value

def records(book, sheet):
    ws = book[sheet]
    rows = ws.iter_rows(values_only=True)
    headers = [str(v).strip() for v in next(rows)]
    for number, values in enumerate(rows, 2):
        row = {key: clean(value) for key, value in zip(headers, values)}
        if any(value not in (None, '') for value in row.values()):
            yield number, row

def required(row, fields, location):
    missing = [key for key in fields if row.get(key) in (None, '')]
    if missing:
        raise ValueError(f'{location}: missing {", ".join(missing)}')

def checked_dates(row, location):
    for key, value in row.items():
        if not (key.endswith('_date') and value not in (None, '')):
            continue
        try:
            row[key] = date.fromisoformat(str(value)).isoformat()
        except ValueError as error:
            raise ValueError(f'{location}: {key} must be a valid yyyy-mm-dd date') from error

def entities(book, sheet, key, fields):
    result = []
    seen = set()
    for number, row in records(book, sheet):
        place = f'{sheet} row {number}'
        checked_dates(row, place)
        required(row, [key], place)
        if row[key] in seen:
            raise ValueError(f'{place}: duplicate {key} {row[key]}')
        seen.add(row[key])
        result.append({field: row.get(field) for field in fields})
    return result

def edges(book, sheet, left, right, known):
    result = []
    for number, row in records(book, sheet):
        place = f'{sheet} row {number}'
        checked_dates(row, place)
        required(row, [left, right], place)
        if row[left] not in known[left] or row[right] not in known[right]:
            raise ValueError(f'{place}: unknown ID in {left}/{right}')
        entry = {left: row[left], right: row[right]}
        for field in ('start_date', 'end_date', 'author_order'):
            if row.get(field) not in (None, ''):
                entry[field] = row[field]
        if entry.get('start_date') and entry.get('end_date') and entry['end_date'] <= entry['start_date']:
            raise ValueError(f'{place}: end_date must be later than start_date')
        result.append(entry)
    return result

def build(input_path, output_path, include_internal_email=False):
    book = load_workbook(input_path, read_only=True, data_only=True)
    fields = {
        'themes': ('Themes', 'theme_id', ['theme_id', 'RAP', 'theme_name', 'color_hex', 'light_hex', 'description']),
        'cases': ('Case studies', 'case_id', ['case_id', 'case_name', 'description', 'status', 'start_date', 'end_date']),
        'teams': ('Teams', 'team_id', ['team_id', 'team_name', 'institution', 'discipline', 'country', 'start_date', 'end_date']),
        'people': ('People', 'person_id', ['person_id', 'full_name', 'role', 'institution', 'team_id', 'email', 'email_visibility', 'website_or_orcid', 'start_date', 'end_date']),
        'publications': ('Publications', 'publication_id', ['publication_id', 'title', 'year', 'first_visible_date', 'doi_or_url', 'abstract_or_note']),
        'keywords': ('Keywords', 'keyword_id', ['keyword_id', 'keyword_label']),
    }
    data = {name: entities(book, *spec) for name, spec in fields.items()}
    known = {key: {item[key] for item in data[name]} for name, (_, key, _) in fields.items()}
    for person in data['people']:
        if person['team_id'] not in known['team_id']:
            raise ValueError(f"People: unknown team_id for {person['person_id']}")
        if not include_internal_email or person['email_visibility'] != 'Internal map':
            person['email'] = None
        person.pop('email_visibility')
    relations = [
        ('themeCase', 'Theme Case', 'theme_id', 'case_id'),
        ('teamTheme', 'Team Theme', 'team_id', 'theme_id'),
        ('teamCase', 'Team Case', 'team_id', 'case_id'),
        ('personCase', 'Person Case', 'person_id', 'case_id'),
        ('personTheme', 'Person Theme', 'person_id', 'theme_id'),
        ('publicationAuthor', 'Publication Author', 'publication_id', 'person_id'),
        ('publicationKeyword', 'Publication Keyword', 'publication_id', 'keyword_id'),
        ('publicationCase', 'Publication Case', 'publication_id', 'case_id'),
        ('publicationTheme', 'Publication Theme', 'publication_id', 'theme_id'),
    ]
    data['links'] = {name: edges(book, sheet, left, right, known) for name, sheet, left, right in relations}
    statuses = []
    for number, row in records(book, 'Publication Status'):
        checked_dates(row, f'Publication Status row {number}')
        required(row, ['publication_id', 'effective_date', 'status'], f'Publication Status row {number}')
        if row['publication_id'] not in known['publication_id']:
            raise ValueError(f'Publication Status row {number}: unknown publication_id')
        statuses.append({k: row[k] for k in ('publication_id', 'effective_date', 'status')})
    data['publicationStatus'] = statuses
    data['schemaVersion'] = 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Built {output_path} from {input_path}')

if __name__ == '__main__':
    if len(sys.argv) not in (3, 4) or (len(sys.argv) == 4 and sys.argv[3] != '--include-internal-email'):
        raise SystemExit(__doc__)
    build(Path(sys.argv[1]), Path(sys.argv[2]), '--include-internal-email' in sys.argv[3:])
